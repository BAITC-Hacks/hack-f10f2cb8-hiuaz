import io
import unittest
from core import align_words, normalize_deadline, parse_text, validate_extraction, export_docx


class DeadlineTests(unittest.TestCase):
    def test_unknown_meeting_date_is_not_today(self):
        self.assertIsNone(normalize_deadline('завтра')[0])
        self.assertIsNone(normalize_deadline('15 октября')[0])
        self.assertIsNone(normalize_deadline(None, '2026-09-23')[0])

    def test_relative_deadlines_and_ambiguous_weekday(self):
        self.assertEqual(normalize_deadline('ертең', '2026-09-23')[0], '2026-09-24')
        self.assertEqual(normalize_deadline('за две недели', '2026-09-23')[0], '2026-10-07')
        self.assertIsNone(normalize_deadline('к пятнице', '2026-09-23')[0])
        self.assertIsNone(normalize_deadline('после совещания', '2026-09-23')[0])

    def test_exact_date_and_invalid_date(self):
        self.assertEqual(normalize_deadline('до 25 сентября 2026 года')[0], '2026-09-25')
        self.assertIsNone(normalize_deadline('31.02.2026')[0])


class AttributionTests(unittest.TestCase):
    def test_word_speaker_overlap(self):
        words = [{'start': 1, 'end': 2, 'text': ' Задача'}, {'start': 3, 'end': 4, 'text': ' выполнена'}]
        turns = [{'start': 0, 'end': 2.5, 'speaker_id': 'A'}, {'start': 0.9, 'end': 1.6, 'speaker_id': 'B'}]
        rows = align_words(words, turns)
        self.assertEqual(rows[0]['speaker_id'], 'A')
        self.assertTrue(rows[0]['speaker_review'])
        self.assertEqual(rows[1]['speaker_id'], 'UNKNOWN')

    def test_absent_assignee_and_fake_evidence(self):
        rows = parse_text('Руководитель: Ерлан должен проверить договор.\nАйдана: Передам ему.')
        raw = {'summary': 'Договор будет проверен.', 'actions': [
            {'action': 'Проверить договор', 'owner': 'Ерлан', 'deadline_raw': None, 'evidence_ids': ['T0001', 'FAKE'], 'note': ''},
            {'action': 'Уточнить детали', 'owner': 'SPEAKER_01', 'deadline_raw': 'завтра', 'evidence_ids': ['FAKE'], 'note': ''}]}
        result = validate_extraction(raw, rows)
        self.assertEqual(result['actions'][0]['owner'], 'Ерлан')
        self.assertEqual(result['actions'][0]['evidence_ids'], ['T0001'])
        self.assertEqual(result['actions'][1]['owner'], '')
        self.assertEqual(result['actions'][1]['due_date'], '')
        self.assertIn('Источник не подтвержден', result['actions'][1]['note'])

    def test_docx_preserves_ambiguity_and_source(self):
        from docx import Document
        rows = parse_text('Айдана: Сравните предложения, срок уточним.')
        action = {'id': 'A001', 'action': 'Сравнить предложения', 'owner': '', 'deadline_raw': '',
                  'due_date': '', 'evidence_ids': ['T0001'], 'note': 'Нужен исполнитель', 'verified': False}
        payload = export_docx('Проверка', None, 'Обсудили закупку.', [action], rows, 'Текст', True)
        doc = Document(io.BytesIO(payload))
        self.assertEqual(len(doc.tables[0].rows), 2)
        self.assertEqual(doc.tables[0].cell(1, 1).text, 'Не указан')
        self.assertIn('T0001', doc.tables[0].cell(1, 3).text)
        self.assertTrue(any('фрагмент' in p.text for p in doc.paragraphs))


if __name__ == '__main__':
    unittest.main()
